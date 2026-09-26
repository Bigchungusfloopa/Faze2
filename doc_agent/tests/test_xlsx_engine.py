"""
Comprehensive Generic XLSX Intelligence Engine Test Suite
=========================================================
Tests the dataset-independent XLSX processing engine across multiple structurally
different datasets (Section 25) and verifies all anti-hallucination tests A through F (Section 26).
"""

import os
import openpyxl
import pandas as pd
import pytest
from pathlib import Path

from agent.xlsx_engine import (
    XLSXIntelligenceEngine,
    XLSXStatus,
    ColumnDataType,
    detect_column_data_type,
    inspect_workbook,
)


# ============================================================================
# FIXTURES: MULTIPLE STRUCTURALLY DIFFERENT DATASETS (Section 25)
# ============================================================================

@pytest.fixture
def student_xlsx(tmp_path) -> Path:
    """Dataset A: Student information (Section 25)."""
    file_path = tmp_path / "students.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Students_2024"
    ws.append(["Student_ID", "Name", "Grade", "Attendance_Pct", "Major"])
    ws.append(["STU001", "Alice Smith", 88.5, 94.2, "Computer Science"])
    ws.append(["STU002", "Bob Jones", 72.0, 81.0, "Mathematics"])
    ws.append(["STU003", "Charlie Brown", 95.0, 98.5, "Computer Science"])
    ws.append(["STU004", "Diana Prince", 64.0, 75.0, "Physics"])
    ws.append(["STU005", "Evan Wright", 82.5, 88.0, "Mathematics"])
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def sales_xlsx(tmp_path) -> Path:
    """Dataset B: Sales transactions (Section 25)."""
    file_path = tmp_path / "sales_data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Transactions"
    ws.append(["Transaction_ID", "Category", "Units_Sold", "Revenue", "Discount"])
    ws.append(["TX101", "Electronics", 15, 7500.0, 0.10])
    ws.append(["TX102", "Furniture", 4, 1200.0, 0.05])
    ws.append(["TX103", "Electronics", 20, 10000.0, 0.15])
    ws.append(["TX104", "Appliances", 8, 3200.0, 0.00])
    ws.append(["TX105", "Furniture", 10, 3000.0, 0.10])
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def weather_xlsx(tmp_path) -> Path:
    """Dataset C: Weather observations (Section 25)."""
    file_path = tmp_path / "weather_metrics.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Observations"
    ws.append(["City", "Temperature_C", "Humidity_Pct", "Wind_Speed_Kmh", "Condition"])
    ws.append(["London", 14.5, 82.0, 15.2, "Rainy"])
    ws.append(["Tokyo", 22.0, 65.0, 8.4, "Clear"])
    ws.append(["New York", 18.2, 70.0, 12.1, "Cloudy"])
    ws.append(["Sydney", 26.5, 55.0, 18.0, "Sunny"])
    ws.append(["Paris", 16.0, 78.0, 10.5, "Overcast"])
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def financial_xlsx(tmp_path) -> Path:
    """Dataset D: Financial portfolio (Section 25)."""
    file_path = tmp_path / "financial_portfolio.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Holdings"
    ws.append(["Ticker", "Asset_Class", "Shares", "Purchase_Price", "Current_Price", "Dividend_Yield"])
    ws.append(["AAPL", "Equity", 50, 150.0, 180.0, 0.006])
    ws.append(["GOOGL", "Equity", 30, 130.0, 145.0, 0.000])
    ws.append(["BND", "Fixed Income", 100, 72.0, 74.0, 0.035])
    ws.append(["MSFT", "Equity", 40, 310.0, 370.0, 0.008])
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def healthcare_xlsx(tmp_path) -> Path:
    """Dataset E: Healthcare/hospital admissions (Section 25)."""
    file_path = tmp_path / "hospital_admissions.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Inpatients"
    ws.append(["Patient_ID", "Age", "Systolic_BP", "Cholesterol", "Department", "Stay_Days"])
    ws.append(["P001", 45, 128, 210, "Cardiology", 4])
    ws.append(["P002", 62, 142, 240, "Cardiology", 7])
    ws.append(["P003", 29, 118, 175, "Neurology", 2])
    ws.append(["P004", 51, 135, 195, "Orthopedics", 5])
    ws.append(["P005", 73, 150, 260, "Cardiology", 9])
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def relational_multisheet_xlsx(tmp_path) -> Path:
    """Dataset F: Multi-sheet relational dataset with shared Customer_ID (Section 14 & 26 Test E)."""
    file_path = tmp_path / "customers_orders.xlsx"
    wb = openpyxl.Workbook()
    
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Customers"
    ws1.append(["Customer_ID", "Customer_Name", "Region"])
    ws1.append(["C101", "Acme Corp", "North America"])
    ws1.append(["C102", "Globex Inc", "Europe"])
    ws1.append(["C103", "Soylent Co", "Asia"])
    
    ws2 = wb.create_sheet(title="Orders")
    ws2.append(["Order_ID", "Customer_ID", "Order_Value", "Status"])
    ws2.append(["ORD-01", "C101", 1500.0, "Shipped"])
    ws2.append(["ORD-02", "C101", 2500.0, "Delivered"])
    ws2.append(["ORD-03", "C102", 4200.0, "Delivered"])
    ws2.append(["ORD-04", "C103", 1100.0, "Pending"])
    
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def unrelated_multisheet_xlsx(tmp_path) -> Path:
    """Dataset with multiple sheets sharing NO common key (Section 14 & 26 Test E)."""
    file_path = tmp_path / "unrelated_sheets.xlsx"
    wb = openpyxl.Workbook()
    
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Sensors"
    ws1.append(["Sensor_Code", "Frequency_Hz", "Calibration_Date"])
    ws1.append(["SENS-1", 440, "2024-01-10"])
    
    ws2 = wb.create_sheet(title="Cafeteria_Menu")
    ws2.append(["Dish_Name", "Calories", "Price_USD"])
    ws2.append(["Salad", 250, 6.50])
    
    wb.save(str(file_path))
    return file_path


@pytest.fixture
def conflicting_multisheet_xlsx(tmp_path) -> Path:
    """Dataset with conflicting data across sheets for the same ID (Section 26 Test F)."""
    file_path = tmp_path / "conflicting_records.xlsx"
    wb = openpyxl.Workbook()
    
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Q1_Status"
    ws1.append(["Student_ID", "Status", "Grade"])
    ws1.append(["STU001", "Active", 88.5])
    ws1.append(["STU002", "Active", 72.0])
    
    ws2 = wb.create_sheet(title="Q2_Status")
    ws2.append(["Student_ID", "Status", "Grade"])
    ws2.append(["STU001", "Suspended", 88.5])  # CONFLICT: Status is Active in Q1 but Suspended in Q2
    ws2.append(["STU002", "Active", 72.0])
    
    wb.save(str(file_path))
    return file_path


# ============================================================================
# 1. DYNAMIC DISCOVERY & SCHEMA EXTRACTION TESTS (Sections 3, 4, 5)
# ============================================================================

def test_dynamic_schema_discovery(student_xlsx):
    """Verifies workbook discovery without any hardcoded names (Sections 3 & 4)."""
    engine = XLSXIntelligenceEngine()
    meta = engine.inspect(str(student_xlsx))
    
    assert meta.sheet_count == 1
    assert "Students_2024" in meta.sheet_names
    schema = meta.sheets["Students_2024"]
    assert schema.row_count == 5
    assert schema.column_count == 5
    
    # Check Section 4 JSON structure compliance
    schema_dict = schema.to_dict()
    assert schema_dict["sheet_name"] == "Students_2024"
    assert schema_dict["row_count"] == 5
    assert schema_dict["column_count"] == 5
    assert len(schema_dict["columns"]) == 5
    for col_info in schema_dict["columns"]:
        assert "name" in col_info
        assert "dtype" in col_info
        assert "null_count" in col_info
        assert "unique_count" in col_info


def test_dynamic_data_type_detection():
    """Verifies detection of all granular data types (Section 5)."""
    s_int = pd.Series([10, 20, 30, 40])
    assert detect_column_data_type(s_int, "Score") == ColumnDataType.INTEGER

    s_float = pd.Series([10.5, 20.2, 30.1])
    assert detect_column_data_type(s_float, "Rate") == ColumnDataType.FLOAT

    s_bool = pd.Series([True, False, True])
    assert detect_column_data_type(s_bool, "Is_Active") == ColumnDataType.BOOLEAN

    s_date = pd.Series(["2024-01-15", "2024-02-20", "2024-03-25"])
    assert detect_column_data_type(s_date, "Joining_Date") in (ColumnDataType.DATE, ColumnDataType.DATETIME)

    s_time = pd.Series(["09:30:00", "14:15:00", "18:45:00"])
    assert detect_column_data_type(s_time, "Shift_Time") == ColumnDataType.TIME

    s_id = pd.Series(["ID_101", "ID_102", "ID_103", "ID_104"])
    assert detect_column_data_type(s_id, "User_ID") == ColumnDataType.ID_LIKE

    s_cat = pd.Series(["North", "South", "East", "West", "North", "South"] * 10)
    assert detect_column_data_type(s_cat, "Region") == ColumnDataType.CATEGORICAL

    s_empty = pd.Series([None, None, None])
    assert detect_column_data_type(s_empty, "Notes") == ColumnDataType.EMPTY


# ============================================================================
# 2. HALLUCINATION TESTS A THROUGH F (Section 26)
# ============================================================================

def test_hallucination_test_a_valid_question(student_xlsx):
    """
    Test A — Valid question:
    Question refers to an existing column.
    Expected: verified (Section 26).
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(student_xlsx))
    
    resp = engine.query("What is the average Attendance_Pct?")
    assert resp.status == XLSXStatus.VERIFIED
    assert len(resp.evidence) > 0
    assert resp.evidence[0]["columns"] == ["Attendance_Pct"]
    assert resp.evidence[0]["operation"] == "mean"
    assert resp.evidence[0]["rows_analyzed"] == 5
    assert resp.evidence[0]["result"] == pytest.approx(87.34, 0.01)


def test_hallucination_test_b_nonexistent_column(student_xlsx):
    """
    Test B — Nonexistent column:
    Question refers to a column not present (e.g. asking for 'revenue' on a student dataset).
    Expected: insufficient_information (Section 26 & Section 18).
    MUST NOT substitute or invent columns.
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(student_xlsx))
    
    resp = engine.query("What is the average revenue?")
    assert resp.status == XLSXStatus.INSUFFICIENT_INFORMATION
    assert "revenue" in resp.answer.lower()
    assert "Insufficient information" in resp.answer
    assert resp.evidence == []


def test_hallucination_test_c_ambiguous_column(student_xlsx):
    """
    Test C — Ambiguous column:
    Multiple columns could answer the question (e.g. 'What is the average?').
    Expected: clarification_required (Section 26 & Section 19).
    DO NOT choose one.
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(student_xlsx))
    
    resp = engine.query("What is the average?")
    assert resp.status == XLSXStatus.CLARIFICATION_REQUIRED
    assert resp.clarification_options is not None
    assert "Grade" in resp.clarification_options
    assert "Attendance_Pct" in resp.clarification_options
    assert "Which column would you like me to calculate" in resp.answer


def test_hallucination_test_d_unsupported_prediction(sales_xlsx):
    """
    Test D — Unsupported prediction:
    Question asks for future prediction or external information not in dataset.
    Expected: insufficient_information (Section 26 & Section 17).
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(sales_xlsx))
    
    resp = engine.query("What will total revenue be next year in 2030?")
    assert resp.status == XLSXStatus.INSUFFICIENT_INFORMATION
    assert "Insufficient Information" in resp.answer
    assert "future projections" in resp.answer.lower()
    assert resp.evidence == []


def test_hallucination_test_e1_multisheet_valid_join(relational_multisheet_xlsx):
    """
    Test E — Multiple sheets (Valid shared key):
    Question requires information from multiple sheets with verified shared key.
    Expected: verified (Section 26 & Section 14).
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(relational_multisheet_xlsx))
    
    resp = engine.query("What is the total order value for each customer across both sheets?")
    assert resp.status == XLSXStatus.VERIFIED
    assert len(resp.evidence) > 0
    assert resp.evidence[0]["join_key"] == "Customer_ID"
    assert resp.evidence[0]["rows_analyzed"] == 4
    assert "Acme Corp" in resp.answer or "C101" in resp.answer


def test_hallucination_test_e2_multisheet_no_shared_key(unrelated_multisheet_xlsx):
    """
    Test E — Multiple sheets (Unrelated tables):
    Do not automatically join unrelated tables.
    Expected: insufficient_information (Section 14 & 26).
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(unrelated_multisheet_xlsx))
    
    resp = engine.query("What is the total price for each sensor across both sheets?")
    assert resp.status == XLSXStatus.INSUFFICIENT_INFORMATION
    assert "Insufficient information to determine a reliable relationship" in resp.answer


def test_hallucination_test_f_conflicting_data(conflicting_multisheet_xlsx):
    """
    Test F — Conflicting data:
    If relevant sheets/sources contain conflicting information for the same entity:
    Expected: conflicting_information (Section 26).
    Do not silently choose one.
    """
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(conflicting_multisheet_xlsx))
    
    resp = engine.query("What is the status of student STU001?")
    assert resp.status == XLSXStatus.CONFLICTING_INFORMATION
    assert "Conflicting information detected" in resp.answer
    assert "Q1_Status" in resp.answer
    assert "Q2_Status" in resp.answer
    assert len(resp.evidence) == 2


# ============================================================================
# 3. STATISTICAL & ADVANCED OPERATIONS ACROSS DATASETS (Sections 8, 9, 10, 11, 12, 16)
# ============================================================================

def test_statistical_operations(financial_xlsx):
    """Verifies count, sum, mean, median, mode, min, max, std, variance, range (Section 8)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(financial_xlsx))

    # Sum
    resp_sum = engine.query("What is the total Shares?")
    assert resp_sum.status == XLSXStatus.VERIFIED
    assert resp_sum.evidence[0]["result"] == 220.0

    # Max
    resp_max = engine.query("What is the maximum Current_Price?")
    assert resp_max.status == XLSXStatus.VERIFIED
    assert resp_max.evidence[0]["result"] == 370.0

    # Min
    resp_min = engine.query("What is the minimum Purchase_Price?")
    assert resp_min.status == XLSXStatus.VERIFIED
    assert resp_min.evidence[0]["result"] == 72.0

    # Range
    resp_range = engine.query("What is the range of Current_Price?")
    assert resp_range.status == XLSXStatus.VERIFIED
    assert resp_range.evidence[0]["result"] == 370.0 - 74.0

    # Standard Deviation
    resp_std = engine.query("What is the standard deviation of Current_Price?")
    assert resp_std.status == XLSXStatus.VERIFIED
    assert resp_std.evidence[0]["result"] > 0


def test_grouping_and_comparison(sales_xlsx):
    """Verifies grouping and category comparison (Sections 10 & 11)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(sales_xlsx))

    # Highest total revenue by category
    resp = engine.query("Which category has the highest total Revenue?")
    assert resp.status == XLSXStatus.VERIFIED
    assert "Electronics" in resp.answer
    assert resp.evidence[0]["operation"] == "groupby_sum"
    assert resp.evidence[0]["result"]["Electronics"] == 17500.0


def test_filtering(healthcare_xlsx):
    """Verifies natural-language filtering (Section 9)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(healthcare_xlsx))

    # Age greater than 50
    resp = engine.query("Show records where Age is greater than 50")
    assert resp.status == XLSXStatus.VERIFIED
    assert resp.evidence[0]["matching_rows"] == 3
    assert "**3** records" in resp.answer


def test_correlation(weather_xlsx):
    """Verifies Pearson correlation between numeric variables (Section 12)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(weather_xlsx))

    resp = engine.query("What is the correlation between Temperature_C and Humidity_Pct?")
    assert resp.status == XLSXStatus.VERIFIED
    assert resp.evidence[0]["operation"] == "pearson_correlation"
    assert isinstance(resp.evidence[0]["result"], float)


def test_cell_lookup(weather_xlsx):
    """Verifies direct cell coordinate resolution (Section 16)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(weather_xlsx))

    resp = engine.query("What is the value in cell B2?")
    assert resp.status == XLSXStatus.VERIFIED
    assert resp.evidence[0]["cell"] == "B2"
    assert "14.5" in resp.answer


def test_metadata_questions(weather_xlsx):
    """Verifies generic dataset-level questions (Section 7)."""
    engine = XLSXIntelligenceEngine()
    engine.inspect(str(weather_xlsx))

    # Row count
    resp_rows = engine.query("How many rows are in this dataset?")
    assert resp_rows.status == XLSXStatus.VERIFIED
    assert resp_rows.evidence[0]["result"] == 5

    # Column count
    resp_cols = engine.query("How many columns are there?")
    assert resp_cols.status == XLSXStatus.VERIFIED
    assert resp_cols.evidence[0]["result"] == 5

    # Sheet list
    resp_sheets = engine.query("What sheets are available?")
    assert resp_sheets.status == XLSXStatus.VERIFIED
    assert "Observations" in resp_sheets.answer
