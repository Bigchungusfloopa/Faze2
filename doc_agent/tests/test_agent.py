import os
import sys
from typing import Any

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest

from agent.gemini_client import MockGeminiClient
from agent.rag_agent import DocumentIntelligenceAgent, INSUFFICIENT_EVIDENCE_MSG
from agent.router import DocumentRouter, PipelineType
from tests.generate_test_docs import generate_all


@pytest.fixture(scope="module")
def docs():
    return generate_all()


@pytest.fixture()
def agent():
    return DocumentIntelligenceAgent(llm_client=MockGeminiClient())


# --------------------------------------------------------------- ROUTING

def test_router_detects_text_pdf(docs):
    decision = DocumentRouter().route(docs["policy"])
    assert decision.pipeline == PipelineType.TEXT_PDF


def test_router_detects_table_pdf(docs):
    decision = DocumentRouter().route(docs["table"])
    assert decision.pipeline == PipelineType.TABLE


def test_router_detects_scanned_pdf(docs):
    decision = DocumentRouter().route(docs["scanned"])
    assert decision.pipeline == PipelineType.SCANNED_PDF


def test_router_detects_image(docs):
    decision = DocumentRouter().route(docs["image"])
    assert decision.pipeline == PipelineType.IMAGE


def test_router_rejects_unsupported_extension(tmp_path):
    bogus = tmp_path / "notes.xyz"
    bogus.write_text("hello")
    decision = DocumentRouter().route(str(bogus))
    assert decision.pipeline == PipelineType.UNSUPPORTED


def test_router_handles_missing_file():
    decision = DocumentRouter().route("/no/such/file.pdf")
    assert decision.pipeline == PipelineType.UNSUPPORTED


# ------------------------------------------------------------- INGESTION

def test_ingest_text_pdf_produces_chunks(agent, docs):
    result = agent.ingest(docs["policy"])
    assert result["status"] == "ok"
    assert result["pipeline"] == "text_pdf"
    assert result["chunk_count"] >= 2  # 2 pages minimum


def test_ingest_table_pdf_extracts_structured_rows(agent, docs):
    result = agent.ingest(docs["table"])
    assert result["status"] == "ok"
    assert result["pipeline"] == "table"


def test_ingest_scanned_pdf_uses_vision_pipeline(agent, docs):
    result = agent.ingest(docs["scanned"])
    assert result["status"] == "ok"
    assert result["pipeline"] == "scanned_pdf"
    assert result["chunk_count"] >= 1


def test_ingest_image_uses_vision_pipeline(agent, docs):
    result = agent.ingest(docs["image"])
    assert result["status"] == "ok"
    assert result["pipeline"] == "image"


def test_ingest_csv_and_query(agent, tmp_path):
    csv_file = tmp_path / "inventory.csv"
    csv_file.write_text("Item,Quantity,UnitCost\nLaptop,25,1200\nMonitor,50,300\n", encoding="utf-8")
    result = agent.ingest(str(csv_file))
    assert result["status"] == "ok"
    assert result["pipeline"] == "csv"
    assert result["chunk_count"] >= 1
    resp = agent.query("What is the Quantity of Monitor?")
    assert not resp.insufficient_evidence
    assert "50" in resp.answer
    assert any("inventory.csv" in c["filename"] for c in resp.citations)


def test_ingest_excel_and_query(agent, tmp_path):
    import openpyxl
    xlsx_file = tmp_path / "finances.xlsx"
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    if ws is None:
        ws = wb.create_sheet(title="Sales")
    else:
        ws.title = "Sales"
    ws.append(["Region", "Revenue"])
    ws.append(["Global", "$950,000"])
    wb.save(str(xlsx_file))
    result = agent.ingest(str(xlsx_file))
    assert result["status"] == "ok"
    assert result["pipeline"] == "excel"
    assert result["chunk_count"] >= 1
    resp = agent.query("What is the Global Revenue?")
    assert not resp.insufficient_evidence
    assert "$950,000" in resp.answer or "950,000" in resp.answer
    assert any("finances.xlsx" in c["filename"] for c in resp.citations)


def test_ingest_docx_and_query(agent, tmp_path):
    import docx
    docx_file = tmp_path / "policy_brief.docx"
    doc = docx.Document()
    doc.add_heading("Travel Policy 2026", level=1)
    doc.add_paragraph("Per diem allowance for domestic travel is $75 per day.")
    doc.save(str(docx_file))
    result = agent.ingest(str(docx_file))
    assert result["status"] == "ok"
    assert result["pipeline"] == "docx"
    resp = agent.query("What is the per diem allowance for domestic travel?")
    assert not resp.insufficient_evidence
    assert "$75" in resp.answer or "75" in resp.answer
    assert any("policy_brief.docx" in c["filename"] for c in resp.citations)


def test_ingest_text_doc_and_query(agent, tmp_path):
    txt_file = tmp_path / "system_manual.md"
    txt_file.write_text("# Server Ports\n\nThe primary backend server operates on port 8080.", encoding="utf-8")
    result = agent.ingest(str(txt_file))
    assert result["status"] == "ok"
    assert result["pipeline"] == "text_doc"
    resp = agent.query("What port does the backend server operate on?")
    assert not resp.insufficient_evidence
    assert "8080" in resp.answer
    assert any("system_manual.md" in c["filename"] for c in resp.citations)


def test_ingest_unsupported_file_is_rejected_cleanly(agent, tmp_path):
    bogus = tmp_path / "data.exe"
    bogus.write_bytes(b"\x00\x01")
    result = agent.ingest(str(bogus))
    assert result["status"] == "rejected"


# ----------------------------------------------------- GROUNDED RETRIEVAL

def test_answer_from_text_pdf_has_citation(agent, docs):
    agent.ingest(docs["policy"])
    resp = agent.query("How many days per week can employees work remotely?")
    assert not resp.insufficient_evidence
    assert resp.grounded
    assert any("policy.pdf" in c["filename"] for c in resp.citations)
    assert "3 days" in resp.answer or "three" in resp.answer.lower()


def test_answer_from_table_pdf_has_citation(agent, docs):
    agent.ingest(docs["table"])
    resp = agent.query("What is the engineering department's annual budget?")
    assert not resp.insufficient_evidence
    assert any(c["filename"] == "salary_table.pdf" for c in resp.citations)
    assert "6,200,000" in resp.answer


def test_answer_from_scanned_pdf_via_ocr(agent, docs):
    agent.ingest(docs["scanned"])
    resp = agent.query("Which parking garage is closed and when?")
    assert not resp.insufficient_evidence
    assert "garage" in resp.answer.lower() or "Garage" in resp.answer
    assert any(c["chunk_type"] == "ocr" for c in resp.citations)


def test_answer_from_image_pipeline(agent, docs):
    agent.ingest(docs["image"])
    resp = agent.query("What company name appears in the logo image?")
    assert not resp.insufficient_evidence
    assert "NORTHWIND" in resp.answer.upper()


def test_multi_source_query_combines_documents(agent, docs):
    """The core 'multi-source' requirement: a single question whose answer
    draws from two different ingested documents should cite both."""
    agent.ingest(docs["policy"])
    agent.ingest(docs["table"])
    resp = agent.query("Tell me about remote work days and the engineering budget.")
    filenames = {c["filename"] for c in resp.citations}
    assert "policy.pdf" in filenames
    assert "salary_table.pdf" in filenames


# ------------------------------------------------------- ANTI-FABRICATION

def test_refuses_when_topic_not_in_any_document(agent, docs):
    agent.ingest(docs["policy"])
    resp = agent.query("What is the capital of France?")
    assert resp.insufficient_evidence
    assert resp.answer == INSUFFICIENT_EVIDENCE_MSG
    assert resp.citations == []


def test_refuses_on_empty_corpus(agent):
    resp = agent.query("Anything at all?")
    assert resp.insufficient_evidence


def test_does_not_fabricate_numbers_not_in_docs(agent, docs):
    agent.ingest(docs["table"])
    resp = agent.query("What is the marketing department's headcount?")
    # answer must ground to the actual ingested value (9), never invent one
    assert not resp.insufficient_evidence
    assert "9" in resp.answer


# --------------------------------------------------------- CONVERSATIONAL

def test_follow_up_question_uses_history(agent, docs):
    agent.ingest(docs["policy"])
    first = agent.query("How many remote days are employees allowed?")
    assert not first.insufficient_evidence

    follow_up = agent.query("Who has to approve it?")  # pronoun "it" -> condensed w/ history
    assert not follow_up.insufficient_evidence
    assert "manager" in follow_up.answer.lower()


def test_session_isolation_between_users(docs):
    """Two different session_ids must not leak conversational context."""
    shared_agent = DocumentIntelligenceAgent(llm_client=MockGeminiClient())
    shared_agent.ingest(docs["policy"])

    shared_agent.query("How many remote days are allowed?", session_id="user_a")
    assert shared_agent.sessions["user_b"] == [] if "user_b" in shared_agent.sessions else True
    resp_b = shared_agent.query("Who has to approve it?", session_id="user_b")
    # user_b has no prior turn, so the dangling pronoun "it" has nothing to
    # resolve against beyond the bare question -- retrieval may still find
    # something via keyword "approve", but session state must be separate
    assert len(shared_agent.sessions["user_a"]) == 1
    assert len(shared_agent.sessions["user_b"]) == 1


# ------------------------------------------------------------- VALIDATOR

def test_validator_rejects_ungrounded_answer_shape(agent, docs):
    """Directly exercise the grounding validator with a fabricated-looking
    answer that shares no vocabulary with retrieved context."""
    agent.ingest(docs["policy"])
    fake_chunks = [{"text": "Employees get 3 remote days per week.", "metadata": {}}]
    grounded, overlap = agent._validate_grounding(
        "The company was founded in 1875 by aliens from Mars.", fake_chunks
    )
    assert grounded is False
    assert overlap < 0.15


def test_detects_conflicting_sources(agent, docs):
    """Two documents give different numbers for the same question -- the
    agent must surface the conflict rather than silently picking one."""
    agent.ingest(docs["policy"])            # says 3 days
    agent.ingest(docs["conflicting_policy"])  # says 5 days
    resp = agent.query("How many remote work days are employees allowed?")
    assert not resp.insufficient_evidence
    assert resp.conflicting
    filenames = {c["filename"] for c in resp.citations}
    assert "policy.pdf" in filenames
    assert "updated_policy.pdf" in filenames


def test_validator_accepts_grounded_answer(agent):
    fake_chunks = [{"text": "Employees get 3 remote days per week.", "metadata": {}}]
    grounded, overlap = agent._validate_grounding(
        "Employees are allowed 3 remote days per week.", fake_chunks
    )
    assert grounded is True
