import os
import sys
from pathlib import Path
from typing import Any
import pytest
from PIL import Image

from agent.gemini_client import MockGeminiClient, register_mock_ocr
from agent.rag_agent import DocumentIntelligenceAgent
from agent.router import DocumentRouter, PipelineType


@pytest.fixture()
def agent():
    return DocumentIntelligenceAgent(llm_client=MockGeminiClient())


# =====================================================================
# 1. POWERPOINT PRESENTATIONS (.pptx, .ppt)
# =====================================================================

def test_presentation_pptx_pipeline(tmp_path, agent):
    from pptx import Presentation

    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # title slide
    slide = prs.slides.add_slide(slide_layout)
    title: Any = slide.shapes.title
    subtitle: Any = slide.placeholders[1]

    if title is not None:
        title.text = "Quarterly Business Review Q3"
    if subtitle is not None:
        subtitle.text = "Revenue reached 15 million dollars with 40% growth."

    # Add second slide with bullet points
    bullet_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_layout)
    slide2_title: Any = slide2.shapes.title
    if slide2_title is not None:
        slide2_title.text = "Product Milestones"
    body_shape: Any = slide2.shapes.placeholders[1]
    tf: Any = getattr(body_shape, "text_frame", None)
    if tf is not None:
        tf.text = "Launched DocAgent Enterprise edition."
        p2 = tf.add_paragraph()
        p2.text = "Integrated PowerPoint and multi-format extraction."

    pptx_path = tmp_path / "review.pptx"
    prs.save(str(pptx_path))

    # Test Router
    decision = DocumentRouter().route(str(pptx_path))
    assert decision.pipeline == PipelineType.PRESENTATION

    # Test Ingestion
    res = agent.ingest(str(pptx_path))
    assert res["status"] == "ok"
    assert res["chunk_count"] >= 2
    assert res["pipeline"] == "presentation"

    # Test Query
    query_resp = agent.query("What was the quarterly revenue growth in review?")
    assert not query_resp.insufficient_evidence
    assert any(c["chunk_type"] == "presentation" for c in query_resp.citations)


def test_legacy_ppt_fallback(tmp_path, agent):
    ppt_path = tmp_path / "legacy_slides.ppt"
    ppt_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 binary header ... Project Alpha Launch Status Report 2026 ...")

    decision = DocumentRouter().route(str(ppt_path))
    assert decision.pipeline == PipelineType.PRESENTATION

    res = agent.ingest(str(ppt_path))
    assert res["status"] == "ok"
    assert res["chunk_count"] >= 1


# =====================================================================
# 2. WORD DOCUMENTS (.docs, .docx, .doc)
# =====================================================================

def test_word_docx_and_docs_extension(tmp_path, agent):
    import docx

    doc = docx.Document()
    doc.add_heading("Employee Benefits Handbook", level=1)
    doc.add_paragraph("Employees are eligible for 25 days paid vacation annually.")
    
    # Save as .docx
    docx_path = tmp_path / "handbook.docx"
    doc.save(str(docx_path))

    decision_docx = DocumentRouter().route(str(docx_path))
    assert decision_docx.pipeline == PipelineType.DOCX

    # Save same document with .docs extension
    docs_path = tmp_path / "handbook.docs"
    doc.save(str(docs_path))

    decision_docs = DocumentRouter().route(str(docs_path))
    assert decision_docs.pipeline == PipelineType.DOCX

    res = agent.ingest(str(docs_path))
    assert res["status"] == "ok"
    assert res["pipeline"] == "docx"

    query_resp = agent.query("How many days paid vacation do employees get?")
    assert not query_resp.insufficient_evidence
    assert "25" in query_resp.answer


def test_legacy_doc_fallback(tmp_path, agent):
    doc_path = tmp_path / "old_memo.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 legacy word doc Confidential Project Phoenix Memo 2026")

    decision = DocumentRouter().route(str(doc_path))
    assert decision.pipeline == PipelineType.DOCX

    res = agent.ingest(str(doc_path))
    assert res["status"] == "ok"
    assert res["chunk_count"] >= 1


# =====================================================================
# 3. PHOTO / IMAGE FORMATS (.jpeg, .jpg, .png, .webp, .bmp, .tiff, .gif)
# =====================================================================

@pytest.mark.parametrize("ext", [".jpeg", ".jpg", ".png", ".webp", ".bmp", ".gif", ".tiff"])
def test_all_photo_image_formats_routing_and_ingest(tmp_path, ext, agent):
    img_path = tmp_path / f"sample_photo{ext}"
    img = Image.new("RGB", (200, 100), color=(50, 100, 150))
    img.save(str(img_path))

    with open(img_path, "rb") as f:
        img_bytes = f.read()

    transcribed_text = f"Photo showing high precision robotics in warehouse {ext}."
    register_mock_ocr(img_bytes, transcribed_text)

    decision = DocumentRouter().route(str(img_path))
    assert decision.pipeline in (PipelineType.IMAGE, PipelineType.OCR)

    res = agent.ingest(str(img_path))
    assert res["status"] == "ok"
    assert res["chunk_count"] >= 1
    assert "ocr_metadata" in res


# =====================================================================
# 4. EXCEL AND TABULAR FORMATS (.xlsx, .xls, .xlsm, .csv, .tsv)
# =====================================================================

def test_excel_formats_routing_and_engine(tmp_path, agent):
    import openpyxl

    wb = openpyxl.Workbook()
    ws: Any = wb.active
    if ws is None:
        ws = wb.create_sheet(title="Sales")
    else:
        ws.title = "Sales"
    ws.append(["Region", "Revenue", "Units"])
    ws.append(["North", 45000, 120])
    ws.append(["South", 38000, 95])

    for ext in [".xlsx", ".xlsm"]:
        p = tmp_path / f"quarterly_sales{ext}"
        wb.save(str(p))

        decision = DocumentRouter().route(str(p))
        assert decision.pipeline == PipelineType.EXCEL

        res = agent.ingest(str(p))
        assert res["status"] == "ok"
        assert "xlsx_metadata" in res
        assert "Sales" in res["xlsx_metadata"]["sheets"]


def test_csv_and_tsv_routing_and_engine(tmp_path, agent):
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("Item,Quantity,Price\nWidgetA,10,15.5\nWidgetB,20,30.0\n")

    decision_csv = DocumentRouter().route(str(csv_path))
    assert decision_csv.pipeline == PipelineType.CSV

    res_csv = agent.ingest(str(csv_path))
    assert res_csv["status"] == "ok"
    assert "xlsx_metadata" in res_csv

    tsv_path = tmp_path / "data.tsv"
    tsv_path.write_text("Item\tQuantity\tPrice\nWidgetA\t10\t15.5\nWidgetB\t20\t30.0\n")

    decision_tsv = DocumentRouter().route(str(tsv_path))
    assert decision_tsv.pipeline == PipelineType.CSV

    res_tsv = agent.ingest(str(tsv_path))
    assert res_tsv["status"] == "ok"
